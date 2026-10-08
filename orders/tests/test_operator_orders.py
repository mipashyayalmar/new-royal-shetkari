"""
Orders through Zomato, Swiggy or Uber Eats (P22).

Since 1 January 2022 an e-commerce operator pays the GST on restaurant
service supplied through it (CGST Act, section 9(5)): the app charges the
guest GST and issues the tax invoice. The restaurant's own bill carries no
GST, and it reports the orders' value in GSTR-1 Table 14 against the app's
GSTIN (and in GSTR-3B 3.1.1(ii)).

Before: those orders were totalled with GST at the dish's rate and counted
in the B2CS sheet, so a restaurant filing from Rasova's export paid GST the
app had already paid.

Run: python manage.py test orders.tests.test_operator_orders
"""
import io
import json
from decimal import Decimal as D
from unittest.mock import patch

import openpyxl
from django.test import TestCase
from django.urls import reverse
from hypothesis import given
from hypothesis.extra.django import TestCase as HypothesisTestCase

from accounts.models import User
from core.utils import get_business_date
from django.utils import timezone
from menu.models import MenuCategory, MenuItem
from orders.models import Order
from orders.services.aggregator_webhook import signed_headers
from orders.services.bill_layout import PLAIN_BILL, bill_layout
from orders.services.tax_engine import COMPOSITION, OPERATOR, REGULAR, UNREGISTERED, Line, compute
from orders.tests.money_scenarios import OUTLET_CONFIGS, build_world, create_orders
from orders.tests.test_totals_properties import REAL, bills
from reports.services.export_services import generate_gstr1_excel
from setup.models import AggregatorConfig
from tenants.models import Outlet, SAMPLE_GSTIN, Tenant

SECRET = "operator-orders-secret"


class EngineTest(TestCase):
    def test_an_operators_order_collects_no_gst(self):
        bill = compute([Line(amount=D("200"), rate=D("5")), Line(amount=D("20"), rate=D("5"), charge="parcel")],
                       gst_paid_by_operator=True)
        self.assertEqual((bill.scheme, bill.gst, bill.grand_total), (OPERATOR, D("0"), D("220")))
        self.assertEqual(bill.rows, ())

    def test_prices_marked_as_including_gst_cost_their_menu_price(self):
        bill = compute([Line(amount=D("105"), rate=D("5"))], prices_include_tax=True, gst_paid_by_operator=True)
        self.assertEqual((bill.gst, bill.grand_total), (D("0"), D("105")))

    def test_composition_and_unregistered_keep_their_own_scheme(self):
        lines = [Line(amount=D("100"), rate=D("5"))]
        self.assertEqual(compute(lines, composition=True, gst_paid_by_operator=True).scheme, COMPOSITION)
        self.assertEqual(compute(lines, gst_registered=False, gst_paid_by_operator=True).scheme, UNREGISTERED)
        self.assertEqual(compute(lines).scheme, REGULAR)


class _World(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(name="App Orders Cafe", slug="app-orders")
        self.outlet = Outlet.objects.create(tenant=self.tenant, name="Main", gst_no=SAMPLE_GSTIN)
        self.config = AggregatorConfig.objects.create(
            tenant=self.tenant, outlet=self.outlet, zomato_enabled=True,
            zomato_webhook_secret=SECRET, auto_accept_orders=False, zomato_gstin="29AABCZ5678K1ZQ",
        )
        category = MenuCategory.objects.create(tenant=self.tenant, outlet=self.outlet, name="Mains")
        self.biryani = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=category,
                                               name="Biryani", price=D("250.00"), gst_percentage=D("5"))
        self.cashier = User.objects.create_user(username="app_cashier", password="pw", role="cashier",
                                                tenant=self.tenant, outlet=self.outlet)
        patcher = patch("orders.api.is_ip_allowed", return_value=True)
        patcher.start()
        self.addCleanup(patcher.stop)

    def webhook_order(self, order_id="ZOM-1", qty=2):
        body = json.dumps({"tenant_id": self.tenant.id, "outlet_id": self.outlet.id, "source": "zomato",
                           "aggregator_order_id": order_id,
                           "items": [{"menu_item_id": self.biryani.id, "quantity": qty}]})
        resp = self.client.post(reverse("api-ingest-order"), data=body, content_type="application/json",
                                headers=signed_headers(SECRET, body.encode()))
        self.assertEqual(resp.status_code, 200, resp.content)
        return Order.objects.get(id=resp.json()["order_id"])

    def pos_order(self, source, aggregator_id=""):
        self.client.force_login(self.cashier)
        resp = self.client.post(reverse("create-order"), content_type="application/json", data=json.dumps({
            "cart": [{"id": self.biryani.id, "quantity": 2}], "source": source, "aggregator_id": aggregator_id,
        }))
        self.assertEqual(resp.status_code, 200, resp.content)
        return Order.objects.get(id=resp.json()["order_id"])


class AppOrdersCarryNoGstTest(_World):
    def test_a_webhook_order_from_zomato(self):
        order = self.webhook_order()
        self.assertEqual((order.gst_scheme, order.gst_total, order.grand_total), (OPERATOR, D("0"), D("500")))
        self.assertEqual(order.payments.get().amount, D("500"))

    def test_a_zomato_order_typed_in_at_the_pos(self):
        order = self.pos_order("zomato", "ZOM-77")
        self.assertEqual((order.gst_scheme, order.gst_total, order.grand_total), (OPERATOR, D("0"), D("500")))

    def test_a_takeaway_still_carries_gst(self):
        order = self.pos_order("takeaway")
        self.assertEqual((order.gst_scheme, order.gst_total, order.grand_total), (REGULAR, D("25.00"), D("525")))

    def test_the_bill_says_the_app_pays_the_gst(self):
        layout = bill_layout(self.webhook_order())
        self.assertEqual(layout.title, PLAIN_BILL)
        self.assertEqual(layout.statement, "GST on this order is paid by Zomato (CGST Act, section 9(5))")
        self.assertEqual([row.label for row in layout.rows], ["Subtotal"])
        self.assertEqual(layout.included, ())

    def test_an_app_order_billed_before_the_record_keeps_its_gst(self):
        order = self.webhook_order()
        Order.objects.filter(pk=order.pk).update(tax_summary=None, gst_total=D("25.00"), grand_total=D("525"))
        order = Order.objects.get(pk=order.pk)
        self.assertEqual(sum(row.tax for row in order.tax_rows()), D("25.00"))


class Gstr1Test(_World):
    def workbook(self):
        today = get_business_date(timezone.now(), self.outlet)
        return openpyxl.load_workbook(io.BytesIO(generate_gstr1_excel(self.tenant, self.outlet, today, today)))

    def table14(self, book):
        return [row for row in book["GSTR-1 Table 14 (App orders)"].iter_rows(min_row=5, values_only=True)
                if row[0] == "Liable to pay tax u/s 9(5)"]

    def test_app_orders_are_in_table_14_not_b2cs(self):
        self.webhook_order("ZOM-1", qty=2)
        self.webhook_order("ZOM-2", qty=1)
        book = self.workbook()
        b2cs = [row for row in book["GSTR-1 B2CS"].iter_rows(min_row=5, values_only=True) if row[0] == "OE"]
        self.assertEqual(b2cs, [])
        self.assertEqual(self.table14(book), [
            ("Liable to pay tax u/s 9(5)", "29AABCZ5678K1ZQ", "Zomato", 750, 0, 0, 0, 0),
        ])

    def test_a_missing_gstin_says_where_to_add_it(self):
        self.config.zomato_gstin = ""
        self.config.save(update_fields=["zomato_gstin"])
        self.webhook_order()
        [row] = self.table14(self.workbook())
        self.assertEqual(row[1], "Not set: add it in Aggregator settings")


class AggregatorSettingsTest(_World):
    def setUp(self):
        super().setUp()
        self.owner = User.objects.create_user(username="app_owner", password="pw", role="owner",
                                              tenant=self.tenant, outlet=self.outlet)
        self.client.force_login(self.owner)

    def save(self, **gstins):
        return self.client.post(reverse("setup_aggregators"), {"zomato_enabled": "on", **gstins}, follow=True)

    def test_gstins_are_saved_cleaned_up(self):
        self.save(zomato_gstin=" 29aabcz5678k1zq ", swiggy_gstin="29AABCS1234D1CZ", uber_eats_gstin="")
        self.config.refresh_from_db()
        self.assertEqual((self.config.zomato_gstin, self.config.swiggy_gstin, self.config.uber_eats_gstin),
                         ("29AABCZ5678K1ZQ", "29AABCS1234D1CZ", ""))

    def test_a_wrong_gstin_is_refused_and_the_old_one_stays(self):
        page = self.save(zomato_gstin="ZOMATO").content.decode()
        self.assertIn("ZOMATO is not a valid GSTIN", page)
        self.config.refresh_from_db()
        self.assertEqual(self.config.zomato_gstin, "29AABCZ5678K1ZQ")


class AppOrdersBillLikeCompositionTest(HypothesisTestCase):
    """Any bill, as an app order at a registered outlet, costs exactly what
    the same bill costs on the composition scheme, where GST is off too."""

    MONEY = ("sub", "gst", "disc", "grand", "ro", "cgst", "sgst", "pgst", "bd")

    @classmethod
    def setUpTestData(cls):
        cls.world = build_world()

    @REAL
    @given(bills(cfgs=[i for i, (_, composition) in enumerate(OUTLET_CONFIGS) if not composition]))
    def test_app_orders_bill_like_the_composition_scheme(self, spec):
        from orders.tests.money_scenarios import snapshot
        inclusive, _ = OUTLET_CONFIGS[spec["cfg"]]
        on_composition = {**spec, "cfg": OUTLET_CONFIGS.index((inclusive, True))}
        app, composition = create_orders(self.world, [spec, on_composition])
        Order.objects.filter(pk=app.pk).update(source="swiggy")
        app.source = "swiggy"
        app.recalculate_totals()
        composition.recalculate_totals()
        mine, theirs = snapshot(app, 0, spec), snapshot(composition, 0, on_composition)
        self.assertEqual({k: v for k, v in mine.items() if k in self.MONEY},
                         {k: v for k, v in theirs.items() if k in self.MONEY})
        self.assertEqual(app.gst_scheme, OPERATOR)
