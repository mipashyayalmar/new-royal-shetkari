"""
Reports and returns read each bill's own tax record, so they always agree
with the bills:

  * GSTR-1 counts each dish at the rate it was sold at, not the menu's rate
    today (P2), includes the parcel charge (P17) and adds up to exactly the
    bills' GST
  * the tax inspection screen shows the real taxable value at each rate, and
    totals with their paise (P9)
  * the orders CSV and the printed receipt show Indian time, not UTC (P8)

Run: python manage.py test reports.tests.test_tax_reports
"""
import base64
import csv
import datetime as dt
import io
from decimal import Decimal as D

import openpyxl
from django.test import Client, TestCase
from django.utils import timezone

from core.utils import get_business_date

from accounts.models import User
from menu.models import MenuCategory, MenuItem
from orders.models import Order, OrderItem, Payment
from printing.views import _build_receipt_b64
from reports.services.export_services import generate_gstr1_excel, generate_orders_csv
from setup.models import PaymentConfig
from tenants.models import Outlet, Tenant


class Base(TestCase):

    def setUp(self):
        self.tenant = Tenant.objects.create(name="Tax Reports Test")
        self.outlet = Outlet.objects.create(tenant=self.tenant, name="Main", gst_no="29ABCDE1234F1Z5")
        PaymentConfig.objects.create(tenant=self.tenant, outlet=self.outlet, cash_enabled=True)
        self.owner = User.objects.create_user(username="tax_reports_owner", password="x",
                                              tenant=self.tenant, outlet=self.outlet, role="owner")
        category = MenuCategory.objects.create(tenant=self.tenant, outlet=self.outlet, name="Food")
        self.dosa = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=category,
                                            name="Dosa", price=D("90"), gst_percentage=D("5"))
        self.sandwich = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=category,
                                                name="Sandwich", price=D("320"), gst_percentage=D("18"))
        self.curd = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=category,
                                            name="Curd", price=D("30"), gst_percentage=D("0"))
        self.client = Client()
        self.client.force_login(self.owner)
        self.today = get_business_date(timezone.now(), self.outlet)   # 6 AM to 6 AM

    def paid_bill(self, *lines, parcel=None):
        """Totalled while open, then paid: the real order of events."""
        order = Order.objects.create(tenant=self.tenant, outlet=self.outlet, created_by=self.owner,
                                     source="counter", status="open")
        for item, qty in lines:
            OrderItem.objects.create(order=order, menu_item=item, quantity=qty, price=item.price,
                                     gst_percentage=item.gst_percentage, total_price=item.price * qty)
        if parcel:
            order.parcel_surcharge = D(parcel)
            order.parcel_gst_rate = self.outlet.parcel_gst_rate       # as toggle_parcel copies it
            order.save(update_fields=["parcel_surcharge", "parcel_gst_rate"])
        order.recalculate_totals()
        Payment.objects.create(order=order, method="cash", amount=order.grand_total)
        order.status = "paid"
        order.save(update_fields=["status"])
        return order

    def b2cs_rows(self):
        book = openpyxl.load_workbook(io.BytesIO(generate_gstr1_excel(self.tenant, self.outlet, self.today, self.today)))
        sheet = book["GSTR-1 B2CS"]
        rows = [row for row in sheet.iter_rows(min_row=5, values_only=True) if row[0] == "OE"]
        total = [row for row in sheet.iter_rows(min_row=5, values_only=True) if row[0] == "TOTAL"][0]
        return {row[2]: {"taxable": D(str(row[3])), "cgst": D(str(row[4])), "sgst": D(str(row[5])),
                         "tax": D(str(row[9]))} for row in rows}, total


class Gstr1Test(Base):

    def test_each_dish_counts_at_the_rate_it_was_sold_at(self):
        self.paid_bill((self.dosa, 2))
        self.dosa.gst_percentage = D("18")          # the menu changes after the sale
        self.dosa.save(update_fields=["gst_percentage"])
        rows, _ = self.b2cs_rows()
        self.assertEqual(set(rows), {5})
        self.assertEqual((rows[5]["taxable"], rows[5]["tax"]), (D("180"), D("9")))

    def test_the_parcel_charge_is_in_its_rates_row(self):
        self.paid_bill((self.dosa, 2), parcel="20")      # food 180 + parcel 20, both at 5%
        rows, _ = self.b2cs_rows()
        self.assertEqual((rows[5]["taxable"], rows[5]["tax"]), (D("200"), D("10")))

    def test_the_return_adds_up_to_the_bills(self):
        bills = [
            self.paid_bill((self.dosa, 1), (self.sandwich, 1)),
            self.paid_bill((self.dosa, 3), (self.curd, 1), parcel="10"),
            self.paid_bill((self.sandwich, 3)),
        ]
        _, total = self.b2cs_rows()
        self.assertEqual(D(str(total[9])), sum(b.gst_total for b in bills))
        self.assertEqual(D(str(total[4])), sum(b.cgst_total for b in bills))
        self.assertEqual(D(str(total[5])), sum(b.sgst_total for b in bills))

    def test_summing_bills_does_not_fetch_their_lines(self):
        from reports.services.tax_totals import rate_totals
        for _ in range(8):
            self.paid_bill((self.dosa, 1), (self.sandwich, 1))
        # one query for bills with a record, one for older bills (none here)
        with self.assertNumQueries(2):
            totals = rate_totals(Order.objects.filter(tenant=self.tenant))
        self.assertEqual(totals[("gst", D("5.00"))]["taxable"], D("720.00"))

    def test_a_bill_from_before_the_tax_record_is_counted_the_same(self):
        order = self.paid_bill((self.dosa, 1), (self.sandwich, 1))
        with_record, _ = self.b2cs_rows()
        Order.objects.filter(pk=order.pk).update(tax_summary=None)
        without_record, _ = self.b2cs_rows()
        self.assertEqual(with_record, without_record)


class Gstr1Table8Test(Base):
    """Nil rated and non-GST supplies go in Table 8, not B2CS (P20): B2CS
    is for taxed supplies, and liquor is outside GST altogether."""

    SHEET = "GSTR-1 Table 8 (Nil, non-GST)"
    OURS = "Intra-State supplies to unregistered persons"

    def table8(self):
        book = openpyxl.load_workbook(io.BytesIO(generate_gstr1_excel(self.tenant, self.outlet, self.today, self.today)))
        rows = book[self.SHEET].iter_rows(min_row=5, max_row=8, values_only=True)
        return {row[0]: tuple(D(str(value)) for value in row[1:4]) for row in rows}

    def liquor_bill(self, price, qty, vat_rate):
        beer = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=self.dosa.category,
                                       name="Draught Beer", price=D(price), gst_percentage=D("0"))
        order = self.paid_bill((self.dosa, 2))
        Order.objects.filter(pk=order.pk).update(status="open")
        order.refresh_from_db()
        OrderItem.objects.create(order=order, menu_item=beer, quantity=qty, price=beer.price,
                                 gst_percentage=D("0"), tax_kind="vat", vat_rate=D(vat_rate),
                                 total_price=beer.price * qty)
        order.recalculate_totals()
        Order.objects.filter(pk=order.pk).update(status="paid")
        return order

    def test_zero_percent_dishes_are_nil_rated_not_in_b2cs(self):
        self.paid_bill((self.dosa, 2), (self.curd, 1))
        rows, _ = self.b2cs_rows()
        self.assertEqual(set(rows), {5})
        table = self.table8()
        self.assertEqual(table[self.OURS], (D("30"), D("0"), D("0")))
        self.assertEqual({row for row, values in table.items() if any(values)}, {self.OURS})

    def test_liquor_is_a_non_gst_supply(self):
        self.liquor_bill("320", 2, "0")              # Karnataka: no VAT on liquor
        rows, _ = self.b2cs_rows()
        self.assertEqual(set(rows), {5})
        self.assertEqual(self.table8()[self.OURS], (D("0"), D("0"), D("640")))

    def test_liquor_counts_at_its_value_before_vat(self):
        self.outlet.vat_inclusive = True
        self.outlet.save(update_fields=["vat_inclusive"])
        self.liquor_bill("330", 1, "10")             # 330 includes 30 VAT
        self.assertEqual(self.table8()[self.OURS], (D("0"), D("0"), D("300")))

    def test_a_period_without_either_is_all_zeros(self):
        self.paid_bill((self.dosa, 1))
        self.assertTrue(all(value == 0 for values in self.table8().values() for value in values))


class Gstr1LayoutTest(Base):
    """Every header and figure in the workbook can be read without widening a
    column (P26), on both sheets, under their merged titles."""

    def book(self):
        self.paid_bill((self.dosa, 3), (self.sandwich, 2), (self.curd, 1), parcel="20")
        return openpyxl.load_workbook(io.BytesIO(generate_gstr1_excel(self.tenant, self.outlet, self.today, self.today)))

    def test_nothing_is_cut_off(self):
        for sheet in self.book().worksheets:
            merged = {cell for span in sheet.merged_cells.ranges for cell in span.cells}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None or (cell.row, cell.column) in merged:
                        continue
                    with self.subTest(sheet=sheet.title, cell=cell.coordinate, value=cell.value):
                        self.assertGreaterEqual(sheet.column_dimensions[cell.column_letter].width, len(str(cell.value)))

    def test_the_merged_title_does_not_widen_the_first_column(self):
        for sheet in self.book().worksheets:
            self.assertLess(sheet.column_dimensions["A"].width, len(sheet["A1"].value), sheet.title)


class InspectionTest(Base):

    def test_taxable_value_is_the_value_not_the_tax(self):
        self.paid_bill((self.dosa, 2), (self.sandwich, 1), (self.curd, 1))
        rows = self.client.get("/reports/inspect/").context["gst_by_rate"]
        self.assertEqual(
            [(r["rate"], r["taxable"], r["cgst"], r["sgst"], r["gst"]) for r in rows],
            [(D("0.00"), D("30.00"), D("0.00"), D("0.00"), D("0.00")),
             (D("5.00"), D("180.00"), D("4.50"), D("4.50"), D("9.00")),
             (D("18.00"), D("320.00"), D("28.80"), D("28.80"), D("57.60"))],
        )

    def test_totals_keep_their_paise(self):
        # 450 at 5% = 22.50 (CGST 11.25 + SGST 11.25), 60 at 0%. The old page
        # added CGST and SGST with a filter that dropped the paise: ₹22.
        self.paid_bill((self.dosa, 5), (self.curd, 2))
        response = self.client.get("/reports/inspect/")
        total = response.context["gst_totals"]
        self.assertEqual(total["gst"], D("22.50"))
        self.assertContains(response, "₹22.50")

    def test_a_gst_total_with_odd_paise_is_shown_whole(self):
        dish = MenuItem.objects.create(tenant=self.tenant, outlet=self.outlet, category=self.dosa.category,
                                       name="Thali", price=D("255"), gst_percentage=D("5"))
        self.paid_bill((dish, 2))                         # 510 x 5% = 25.50 = 12.75 + 12.75
        response = self.client.get("/reports/inspect/")
        self.assertEqual(response.context["gst_totals"]["gst"], D("25.50"))
        self.assertContains(response, "₹25.50")
        self.assertNotContains(response, "₹24.00")


class LocalTimeTest(Base):
    """Stored in UTC; shown in the outlet's time (Asia/Kolkata here)."""

    def _bill_at_8pm(self):
        order = self.paid_bill((self.dosa, 1))
        Order.objects.filter(pk=order.pk).update(
            created_at=dt.datetime(2026, 9, 20, 14, 30, tzinfo=dt.timezone.utc))   # 8:00 PM in India
        order.refresh_from_db()
        return order

    def test_the_orders_csv_shows_local_time(self):
        self._bill_at_8pm()
        text = generate_orders_csv(self.tenant, self.outlet, dt.date(2026, 9, 20), dt.date(2026, 9, 20))
        rows = list(csv.reader(io.StringIO(text)))
        header = rows[0]
        row = rows[1]
        self.assertEqual((row[header.index("Date")], row[header.index("Time")]), ("2026-09-20", "20:00:00"))

    def test_the_printed_receipt_shows_local_time(self):
        order = self._bill_at_8pm()
        receipt = base64.b64decode(_build_receipt_b64(order, 48, "full", "cp437")).decode("cp437", "replace")
        self.assertIn("20/09/2026 20:00", receipt)
        self.assertNotIn("14:30", receipt)
